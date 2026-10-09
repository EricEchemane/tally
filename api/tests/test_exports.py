"""Liquidation summary .xlsx and SLSP purchases CSV (SPEC §7, §8)."""
import csv
import io

from openpyxl import load_workbook

import exports
from dbhelpers import add_employee_advance, add_receipt, conn  # noqa: F401  (conn is a fixture)

BF = dict(seller_name="BF Sinangag Express Food Corp", seller_address="Alabang, Muntinlupa", seller_tin="240727673",
          seller_branch="00006", invoice_number="175026", date="2026-04-03", particulars="Rice meals",
          expense_category="MEALS", vat_registered=True, vatable_sales=316.96, vat_amount=38.04, vat_exempt_sales=0,
          zero_rated_sales=0, total=355)


def _csv(text):
    return list(csv.DictReader(io.StringIO(text)))


# ---------- summary .xlsx ----------

def test_summary_xlsx_columns_rows_and_totals(conn):
    add_receipt(conn, 1, "ACCEPTED", corrections={"withholding_tax": 710}, **BF)
    rid = add_receipt(conn, 1, "REJECT", [("DOC_TYPE", "official_receipt", "REJECT"),
                                          ("DATE", "missing", "NEEDS_REVIEW", "DISMISSED")],
                      seller_name="Cebu Hotel", total=560, date="2026-04-05")
    conn.execute("UPDATE finding SET dismissal_reason = 'Date is on the back' WHERE receipt_id = ? AND \"check\" = 'DATE'", (rid,))
    conn.execute("INSERT INTO receipt (liquidation_id, stage, image_path, image_hash) VALUES (1, 'extracting', 'y.jpg', 'h')")

    ws = load_workbook(io.BytesIO(exports.summary_xlsx(conn, 1))).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0][0] == "Liquidation #1"
    assert ["Company", "Bayanihan Demo Corp (TIN 123-456-789-00000)"] == rows[1][:2]
    assert ["Employee", "Juan Dela Cruz"] == rows[2][:2]

    h = next(i for i, r in enumerate(rows) if r[0] == "TIN")
    assert rows[h] == ["TIN", "Date", "Seller name", "Seller address", "Particulars", "Amount", "Withholding tax", "VAT",
                       "VAT-exempt", "Status", "Findings"]
    bf, hotel, pending = rows[h + 1], rows[h + 2], rows[h + 3]
    assert bf[0] == "240-727-673-00006" and bf[1].date().isoformat() == "2026-04-03"
    assert bf[2:10] == ["BF Sinangag Express Food Corp", "Alabang, Muntinlupa", "Rice meals", 355, 7.1, 38.04, 0, "Accepted"]
    assert hotel[9] == "Rejected"
    assert hotel[10] == ("Rejected: DOC_TYPE official_receipt\n"
                         "Dismissed: DATE missing (reason: Date is on the back)")
    assert pending[9] == "Not checked yet"

    t = next(i for i, r in enumerate(rows) if r[0] == "Totals")
    assert [r[:2] for r in rows[t + 1:t + 7]] == [
        ["Cash Advance", 15000], ["Valid Spend", 355], ["Disallowed Amount", 560],
        ["Amount to Return", 14645], ["Excess Spend", 0], ["Claimable input VAT", 38.04],
    ]
    assert ws.cell(h + 2, 6).number_format == "#,##0.00"


def test_summary_xlsx_uses_frozen_totals_once_approved(conn):
    add_receipt(conn, 1, "ACCEPTED", **BF)
    conn.execute("""UPDATE liquidation SET status = 'APPROVED', frozen_totals = '{"advance": 1500000, "valid_spend": 100,
        "disallowed": 0, "amount_to_return": 1499900, "excess_spend": 0, "claimable_input_vat": 0}' WHERE id = 1""")
    ws = load_workbook(io.BytesIO(exports.summary_xlsx(conn, 1))).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert ["Status", "Approved"] == rows[5][:2]
    assert ["Valid Spend", 1] in [r[:2] for r in rows]


# ---------- SLSP CSV ----------

def test_slsp_csv_header():
    assert exports.SLSP_COLUMNS == (
        "month_end,seller_tin,seller_branch_code,registered_name,last_name,first_name,middle_name,address1,address2,"
        "exempt,zero_rated,services,capital_goods,other_goods,taxable_net_of_vat,input_tax,non_creditable_input_tax,"
        "gross_purchases,invoice_no,invoice_date,owner_tin,liquidation_id,receipt_id").split(",")


def test_slsp_csv_rows(conn):
    liq2 = add_employee_advance(conn, "Maria Santos", 3000, "2026-04-01")
    add_receipt(conn, 1, "ACCEPTED", **BF)
    add_receipt(conn, liq2, "NO_INPUT_VAT", [("DOC_TYPE", "not_valid_notice", "NO_INPUT_VAT")],
                **{**BF, "seller_name": "Petron Bacoor", "seller_tin": "000168801", "seller_branch": "123",
                   "expense_category": "TRANSPORT_FUEL", "particulars": "Fuel", "vatable_sales": 1785.71,
                   "vat_amount": 214.29, "total": 2000, "invoice_number": "SI-88", "date": "2026-04-28"})
    add_receipt(conn, 1, "NEEDS_REVIEW", [("SELLER_TIN", "missing", "NEEDS_REVIEW")],
                **{**BF, "seller_name": "Mouthfeel", "seller_tin": None, "seller_branch": None, "vat_registered": False,
                   "vatable_sales": None, "vat_amount": None, "total": 180, "date": "2026-04-10"})
    add_receipt(conn, 1, "REJECT", [("BUYER_DETAILS", "blank", "REJECT")], **{**BF, "date": "2026-04-11"})
    add_receipt(conn, 1, "ACCEPTED", **{**BF, "date": "2026-05-01"})  # next month
    # A corrected Expense Category moves the amount into another SLSP column.
    add_receipt(conn, 1, "ACCEPTED", corrections={"expense_category": "SUPPLIES", "date": "2026-04-30"},
                **{**BF, "seller_tin": "240727673", "invoice_number": "175999"})

    rows = _csv(exports.slsp_csv(conn, "2026-04"))
    assert [(r["registered_name"], r["receipt_id"]) for r in rows] == [
        ("Petron Bacoor", "2"), ("BF Sinangag Express Food Corp", "1"), ("BF Sinangag Express Food Corp", "6"),
        ("Mouthfeel", "3"),
    ]
    petron, bf, bf_supplies, nonvat = rows
    assert bf == {
        "month_end": "04/30/2026", "seller_tin": "240727673", "seller_branch_code": "00006",
        "registered_name": "BF Sinangag Express Food Corp", "last_name": "", "first_name": "", "middle_name": "",
        "address1": "Alabang, Muntinlupa", "address2": "", "exempt": "0.00", "zero_rated": "0.00",
        "services": "316.96", "capital_goods": "0.00", "other_goods": "0.00", "taxable_net_of_vat": "316.96",
        "input_tax": "38.04", "non_creditable_input_tax": "0.00", "gross_purchases": "355.00",
        "invoice_no": "175026", "invoice_date": "04/03/2026", "owner_tin": "123456789",
        "liquidation_id": "1", "receipt_id": "1",
    }
    assert (petron["seller_branch_code"], petron["other_goods"], petron["services"]) == ("00123", "1785.71", "0.00")
    assert (petron["input_tax"], petron["non_creditable_input_tax"], petron["gross_purchases"]) == ("214.29", "214.29", "2000.00")
    assert petron["liquidation_id"] == str(liq2)
    assert (bf_supplies["services"], bf_supplies["other_goods"], bf_supplies["invoice_date"]) == ("0.00", "316.96", "04/30/2026")
    assert (nonvat["seller_tin"], nonvat["exempt"], nonvat["taxable_net_of_vat"], nonvat["input_tax"]) == ("", "180.00", "0.00", "0.00")


def test_slsp_csv_empty_month_has_header_only(conn):
    text = exports.slsp_csv(conn, "2026-02")
    assert text == ",".join(exports.SLSP_COLUMNS) + "\r\n"
