"""Exports (SPEC §7). Owner: Simon.

Interfaces Eric's routes call. Read the DB through db.connect(); amounts in the DB are centavos.
Every number comes from the effective receipt values and serialize.totals_centavos, never from the UI.
"""
import calendar
import csv
import datetime as dt
import io
import json
import sqlite3
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import serialize
from pipeline.derive import slsp_bucket

STATUS_LABELS = {"ACCEPTED": "Accepted", "NO_INPUT_VAT": "No input VAT", "NEEDS_REVIEW": "Needs review",
                 "REJECT": "Rejected", None: "Not checked yet"}
TOTAL_LABELS = [("advance", "Cash Advance"), ("valid_spend", "Valid Spend"), ("disallowed", "Disallowed Amount"),
                ("amount_to_return", "Amount to Return"), ("excess_spend", "Excess Spend"),
                ("claimable_input_vat", "Claimable input VAT")]
MONEY = "#,##0.00"


def _tin(tin: str | None, branch: str | None = None) -> str:
    """'240727673', '00006' -> '240-727-673-00006'."""
    if not tin:
        return ""
    t = f"{tin[:3]}-{tin[3:6]}-{tin[6:9]}"
    return f"{t}-{branch}" if branch else t


def _pesos(centavos: int | None) -> float | None:
    return None if centavos is None else float(Decimal(centavos) / 100)


def _receipts(conn: sqlite3.Connection, where: str, params: tuple) -> list[tuple[sqlite3.Row, dict]]:
    """Receipts with their effective values (corrections over derived), centavos."""
    out = []
    for r in conn.execute(f"SELECT * FROM receipt WHERE {where} ORDER BY id", params):
        eff, _ = serialize.effective(conn, r["id"], json.loads(r["derived"]) if r["derived"] else {})
        out.append((r, eff))
    return out


# ---------- Liquidation summary (.xlsx) ----------

def _findings_text(conn, receipt_id: int) -> str:
    lines = []
    for f in conn.execute('SELECT * FROM finding WHERE receipt_id = ? ORDER BY id', (receipt_id,)):
        if f["state"] == "DISMISSED":
            lines.append(f"Dismissed: {f['message']} (reason: {f['dismissal_reason']})")
        else:
            confirmed = " (confirmed)" if f["state"] == "CONFIRMED" else ""
            lines.append(f"{STATUS_LABELS[f['severity']]}{confirmed}: {f['message']}")
    return "\n".join(lines)


def summary_xlsx(conn: sqlite3.Connection, liquidation_id: int) -> bytes:
    """Liquidation summary: TIN, date, seller name, seller address, particulars, amount,
    withholding tax, VAT, VAT-exempt + status, Findings and a totals block."""
    liq = conn.execute("SELECT * FROM liquidation WHERE id = ?", (liquidation_id,)).fetchone()
    ca = conn.execute("SELECT * FROM cash_advance WHERE id = ?", (liq["cash_advance_id"],)).fetchone()
    employee = conn.execute("SELECT name FROM employee WHERE id = ?", (ca["employee_id"],)).fetchone()["name"]
    company = conn.execute("SELECT name, tin FROM company WHERE id = 1").fetchone()
    totals = serialize.totals_centavos(conn, liquidation_id)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Liquidation {liquidation_id}"
    bold = Font(bold=True)

    ws.append([f"Liquidation #{liquidation_id}"])
    ws["A1"].font = Font(bold=True, size=14)
    header = [
        ("Company", f"{company['name']} (TIN {_tin(company['tin'][:9], company['tin'][9:] or None)})"),
        ("Employee", employee),
        ("Purpose", ca["purpose"]),
        ("Spending Period", f"{ca['released_on']} to {ca['spending_end_on']}"),
        ("Status", liq["status"].capitalize()),
    ]
    if liq["return_note"]:
        header.append(("Return note", liq["return_note"]))
    for label, value in header:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])

    columns = ["TIN", "Date", "Seller name", "Seller address", "Particulars", "Amount", "Withholding tax", "VAT",
               "VAT-exempt", "Status", "Findings"]
    ws.append(columns)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E7EEF7")
    ws.freeze_panes = f"A{header_row + 1}"

    for r, eff in _receipts(conn, "liquidation_id = ?", (liquidation_id,)):
        status = r["status"] if r["stage"] == "done" and r["checks_run"] else None
        date = eff.get("date")
        try:
            date = dt.date.fromisoformat(date) if date else None
        except ValueError:
            pass  # keep the unparsed text
        ws.append([
            _tin(eff.get("seller_tin"), eff.get("seller_branch")), date, eff.get("seller_name"), eff.get("seller_address"),
            eff.get("particulars"), _pesos(eff.get("total")), _pesos(eff.get("withholding_tax")),
            _pesos(eff.get("vat_amount")), _pesos(eff.get("vat_exempt_sales")),
            STATUS_LABELS[status] if r["stage"] != "failed" else "Failed to read", _findings_text(conn, r["id"]),
        ])
        row = ws.max_row
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        for col in (6, 7, 8, 9):
            ws.cell(row, col).number_format = MONEY
        ws.cell(row, 11).alignment = Alignment(wrap_text=True, vertical="top")

    ws.append([])
    ws.append(["Totals"])
    ws.cell(ws.max_row, 1).font = bold
    for key, label in TOTAL_LABELS:
        ws.append([label, _pesos(totals[key])])
        ws.cell(ws.max_row, 2).number_format = MONEY

    for i, width in enumerate([20, 12, 30, 36, 22, 13, 15, 12, 12, 15, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- SLSP-ready purchases CSV ----------

SLSP_COLUMNS = [
    "month_end", "seller_tin", "seller_branch_code", "registered_name", "last_name", "first_name", "middle_name",
    "address1", "address2", "exempt", "zero_rated", "services", "capital_goods", "other_goods", "taxable_net_of_vat",
    "input_tax", "non_creditable_input_tax", "gross_purchases", "invoice_no", "invoice_date", "owner_tin",
    "liquidation_id", "receipt_id",
]


def _amount(centavos: int | None) -> str:
    """Plain decimal with 2 places and no thousands separator, as RELIEF converters expect."""
    return f"{Decimal(centavos or 0) / 100:.2f}"


def _mdy(iso: str | None) -> str:
    try:
        return dt.date.fromisoformat(iso).strftime("%m/%d/%Y") if iso else ""
    except ValueError:
        return ""


def slsp_csv(conn: sqlite3.Connection, month: str) -> str:
    """SLSP-ready purchases CSV for YYYY-MM; non-rejected receipts only. Columns in SPEC §7.

    One row per receipt (the bookkeeper's converter groups by TIN). input_tax is the receipt's whole VAT
    and non_creditable_input_tax repeats it for NO_INPUT_VAT receipts, matching the RELIEF header split
    (input tax = creditable + non-creditable)."""
    year, mon = (int(x) for x in month.split("-"))
    month_end = dt.date(year, mon, calendar.monthrange(year, mon)[1]).strftime("%m/%d/%Y")
    owner_tin = conn.execute("SELECT tin FROM company WHERE id = 1").fetchone()["tin"][:9]

    rows = []
    receipts = _receipts(conn, "stage = 'done' AND checks_run = 1 AND status IS NOT NULL AND status != 'REJECT'", ())
    for r, eff in receipts:
        if not str(eff.get("date") or "").startswith(f"{month}-"):
            continue
        total = eff.get("total") or 0
        zero_rated = eff.get("zero_rated_sales") or 0
        exempt = eff.get("vat_exempt_sales") or 0
        vat = eff.get("vat_amount") or 0
        taxable = eff.get("vatable_sales") or 0
        if eff.get("vat_registered") is False:
            exempt, taxable, vat = total - zero_rated, 0, 0  # a non-VAT seller's purchase carries no VAT
        bucket = slsp_bucket(eff.get("expense_category"), eff.get("particulars"), eff.get("seller_name"))
        branch = "".join(ch for ch in str(eff.get("seller_branch") or "") if ch.isdigit())
        rows.append({
            "month_end": month_end,
            "seller_tin": eff.get("seller_tin") or "",
            "seller_branch_code": branch.zfill(5) if branch else "",
            "registered_name": eff.get("seller_name") or "",
            "last_name": "", "first_name": "", "middle_name": "",
            "address1": eff.get("seller_address") or "", "address2": "",
            "exempt": _amount(exempt),
            "zero_rated": _amount(zero_rated),
            "services": _amount(taxable if bucket == "SERVICES" else 0),
            "capital_goods": _amount(taxable if bucket == "CAPITAL_GOODS" else 0),
            "other_goods": _amount(taxable if bucket == "OTHER_GOODS" else 0),
            "taxable_net_of_vat": _amount(taxable),
            "input_tax": _amount(vat),
            "non_creditable_input_tax": _amount(vat if r["status"] == "NO_INPUT_VAT" else 0),
            "gross_purchases": _amount(total),
            "invoice_no": eff.get("invoice_number") or "",
            "invoice_date": _mdy(eff.get("date")),
            "owner_tin": owner_tin,
            "liquidation_id": r["liquidation_id"],
            "receipt_id": r["id"],
        })
    rows.sort(key=lambda row: (row["seller_tin"] == "", row["seller_tin"], row["registered_name"].upper(),
                               row["invoice_date"][6:] + row["invoice_date"][:5], row["receipt_id"]))

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=SLSP_COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue()
